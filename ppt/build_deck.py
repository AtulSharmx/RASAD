"""Builds ppt/RASAD_SIH26251.pptx by filling the official SIH 2026 template (ppt/template.pptx).

Run: py -3 ppt/build_deck.py   (needs python-pptx and Pillow)
Then export the PDF with PowerPoint (see ppt/README in PLAN.md) and check every slide image.
Every number on the slides comes from the app or docs/ (FACTS, PLACEHOLDERS, VERIFY_REPORT, DEMO_SCRIPT).
"""
import copy
import io
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE.parent / "screenshots"
OUT = HERE / "RASAD_SIH26251.pptx"

# Fill these exactly as registered on the SIH portal; while empty, the template's own labels stay.
TEAM_NAME = "Code ZCrew"
TEAM_ID = "163616"

# template blue + the app's own colours (so diagrams match the screenshots)
NAVY, BLUE, INK, GREY = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x00, 0x70, 0xC0), RGBColor(0x16, 0x21, 0x2B), RGBColor(0x55, 0x5F, 0x6B)
TEAL, GREEN, MAGENTA = RGBColor(0x00, 0x91, 0x7F), RGBColor(0x1D, 0x6F, 0x43), RGBColor(0xB0, 0x42, 0x8F)
AMBER, RED, WHITE = RGBColor(0xA8, 0x6B, 0x00), RGBColor(0xB8, 0x23, 0x2B), RGBColor(0xFF, 0xFF, 0xFF)
TINT, EDGE, PH, PH_EDGE = RGBColor(0xEE, 0xF3, 0xF9), RGBColor(0xC9, 0xD3, 0xDE), RGBColor(0xFF, 0xE5, 0x8A), RGBColor(0xA6, 0x8A, 0x00)
FONT, SYMBOL = "Arial", "Segoe UI Symbol"


# ---------- helpers ----------
def fill_tf(tf, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT):
    """paras: list of str | list of (text, opts) runs | {"runs": [...], "after": pt, "align": ...}."""
    for i, p in enumerate(paras):
        popts = p if isinstance(p, dict) else {}
        runs = p["runs"] if isinstance(p, dict) else [(p, {})] if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = popts.get("align", align)
        if "after" in popts:
            para.space_after = Pt(popts["after"])
        for text, o in runs:
            r = para.add_run()
            r.text = text
            f = r.font
            f.name = o.get("font", FONT)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.underline = o.get("underline", False)
            f.color.rgb = o.get("color", color)
            if o.get("link"):
                r.hyperlink.address = o["link"]


def tb(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, **kw)
    return box


def shp(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, dash=False, paras=None,
        anchor=MSO_ANCHOR.MIDDLE, margin=0.08, radius=None, **kw):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    s.shadow.inherit = False
    if radius is not None:
        s.adjustments[0] = radius
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    if paras:
        kw.setdefault("align", PP_ALIGN.CENTER)
        fill_tf(tf, paras, **kw)
    return s


def label(slide, x, y, w, text):
    """A template pointer, kept word for word as a section label."""
    return tb(slide, x, y, w, 0.34, [[(text, {"bold": True, "color": BLUE})]], size=13.5)


def pic(slide, name, x, y, w, crop=None):
    im = Image.open(SHOTS / name).convert("RGB")
    if crop:
        im = im.crop(crop)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=86, optimize=True)
    buf.seek(0)
    p = slide.shapes.add_picture(buf, Inches(x), Inches(y), width=Inches(w))
    p.line.color.rgb = EDGE
    p.line.width = Pt(1)
    return p


def icon(slide, x, y, d, glyph, fill):
    return shp(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=fill, margin=0,
               paras=[[(glyph, {"font": SYMBOL, "bold": False})]], size=d * 30, color=WHITE)


def tag(slide, x, y, w, text):
    return shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.3, fill=PH, line=PH_EDGE, line_w=0.75, radius=0.3,
               paras=[[(text, {"bold": True})]], size=10.5, color=INK)


def get(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def drop(shape):
    shape._element.getparent().remove(shape._element)


def set_runs(paragraph, text):
    """Replace a paragraph's text but keep the first run's formatting."""
    runs = paragraph.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return runs[0]


def add_value(paragraph, text, size=19):
    """Append a value run after a template label (copied formatting, not bold)."""
    last = paragraph.runs[-1]._r
    new = copy.deepcopy(last)
    new.find(qn("a:t")).text = text
    last.addnext(new)
    run = paragraph.runs[-1]
    run.font.bold = False
    run.font.size = Pt(size)
    run.font.color.rgb = NAVY
    return run


def text_width_in(text, size_pt, font_file="timesbd.ttf"):
    f = ImageFont.truetype(str(Path("C:/Windows/Fonts") / font_file), int(size_pt * 10))
    return f.getlength(text) / 10 / 72


# ---------- build ----------
prs = Presentation(str(HERE / "template.pptx"))
s1, s2, s3, s4, s5, s6, s7 = prs.slides

# drop the instructions slide (the template says to delete it before upload)
ids = prs.slides._sldIdLst
last = ids[6]
prs.part.drop_rel(last.rId)
ids.remove(last)

for s in (s2, s3, s4, s5, s6):  # the "Your Team Name" oval: our name, at a size that fits inside it
    oval = next(sh for sh in s.shapes if sh.name.startswith("Oval"))
    run = set_runs(oval.text_frame.paragraphs[0], TEAM_NAME or "Your Team Name")
    run.font.size, run.font.bold = Pt(11.5), bool(TEAM_NAME)
    oval.text_frame.margin_left = oval.text_frame.margin_right = Inches(0)  # keeps the name on one line
    for p in oval.text_frame.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    if TEAM_ID:
        p = oval.text_frame.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "ID " + TEAM_ID
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

# ----- Slide 1: TITLE PAGE -----
sub = get(s1, "Subtitle 3")
sub.left, sub.top, sub.width, sub.height = Inches(0.36), Inches(0.95), Inches(5.9), Inches(1.35)
tf = sub.text_frame
for p in list(tf.paragraphs):  # the template starts with an empty line before "TITLE PAGE"
    if not p.runs and len(tf.paragraphs) > 1:
        p._p.getparent().remove(p._p)
r = set_runs(tf.paragraphs[0], "RASAD (रसद)")
r.font.size, r.font.color.rgb = Pt(34), NAVY
p = tf.add_paragraph()
p.alignment = tf.paragraphs[0].alignment
r = p.add_run()
r.text = "Decide before the road shuts"
r.font.name, r.font.size, r.font.italic, r.font.bold, r.font.color.rgb = "Times New Roman", Pt(22), True, False, INK

box9 = get(s1, "TextBox 9")
box9.top = Inches(2.45)
fields = box9.text_frame
for para in list(fields.paragraphs):
    if not para.runs:  # the template's empty first line
        para._p.getparent().remove(para._p)
        continue
    # same bullets and weight as the template, but left-aligned and tighter so the long title fits on the slide
    para.alignment = PP_ALIGN.LEFT
    para.line_spacing = 1.45
    for r in para.runs:
        r.font.size = Pt(21)
    t = para.text.strip()
    if t.startswith("Problem Statement ID"):
        add_value(para, " SIH26251")
    elif t.startswith("Problem Statement Title"):
        add_value(para, " Indian Army: Predictive Logistics & Forward Supply Chain")
    elif t.startswith("Theme"):
        add_value(para, " Transportation & Logistics")
    elif t.startswith("PS Category"):
        set_runs(para, "PS Category- ")
        add_value(para, "Software")
    elif t.startswith("Team ID") and TEAM_ID:
        add_value(para, " " + TEAM_ID)
    elif t.startswith("Team Name") and TEAM_NAME:
        set_runs(para, "Team Name- ")
        add_value(para, TEAM_NAME)

# ----- Slide 2: IDEA TITLE -----
drop(get(s2, "TextBox 8"))
title = "RASAD: DECIDE BEFORE THE ROAD SHUTS"
size = 32 if text_width_in(title, 32) <= 8.4 else 28 if text_width_in(title, 28) <= 8.4 else None
t = get(s2, "Title 1").text_frame
if size:
    r = set_runs(t.paragraphs[0], title)
else:
    r = set_runs(t.paragraphs[0], "RASAD (रसद)")
    size = 36
r.font.size = Pt(size)

tb(s2, 0.45, 1.22, 12.4, 0.5,
   [[("❖ ", {"font": SYMBOL}), ("Proposed Solution (Describe your Idea/Solution/Prototype)", {"underline": True})]],
   size=22, bold=True, color=NAVY)

label(s2, 0.45, 1.95, 6.0, "Detailed explanation of the proposed solution")
steps = [("Forecast", "need", NAVY), ("Last safe", "road day", BLUE), ("Mixed-fleet", "plan", TEAL), ("Commander", "approves", GREEN)]
for i, (a, b, col) in enumerate(steps):
    kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
    shp(s2, kind, 0.45 + i * 1.47, 2.32, 1.62, 0.8, fill=col, margin=0.0, radius=0.28,  # shallow point = room for text
        paras=[a, b], size=12.5, bold=True, color=WHITE)

label(s2, 0.45, 3.32, 6.0, "How it addresses the problem")
tb(s2, 0.45, 3.66, 6.0, 0.62, [[("Joins ", {}), ("10 separate data streams", {"bold": True}), (" into one decision picture", {})]],
   size=17)

label(s2, 0.45, 4.38, 6.0, "Innovation and uniqueness of the solution")
chips = [("Closure probability window", BLUE), ("Drones as a real mode", MAGENTA),
         ("Drone basing advisor", MAGENTA), ("Offline, on-premise", GREEN)]
for i, (text, col) in enumerate(chips):
    x, y = 0.45 + (i % 2) * 3.05, 4.74 + (i // 2) * 0.68
    shp(s2, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 2.95, 0.56, fill=TINT, radius=0.25, margin=0.42,
        paras=[[(text, {"bold": True})]], size=14, color=NAVY, align=PP_ALIGN.LEFT)
    shp(s2, MSO_SHAPE.OVAL, x + 0.14, y + 0.18, 0.2, 0.2, fill=col)

pic(s2, "01-dashboard.png", 6.75, 1.95, 6.15)
tb(s2, 6.75, 5.47, 6.15, 0.4,
   [[("Working prototype", {"bold": True, "color": NAVY}), (" · military data simulated · terrain, highways, weather real", {})]],
   size=12, color=GREY, align=PP_ALIGN.LEFT)

# ----- Slide 3: TECHNICAL APPROACH -----
drop(get(s3, "TextBox 8"))
label(s3, 0.4, 1.22, 12.5, "Technologies to be used (e.g. programming languages, frameworks, hardware)")
stack = [("AI / ML", "LightGBM quantile forecast (P10 / P50 / P90)"), ("Optimisation", "Google OR-Tools\nCP-SAT solver"),
         ("GIS", "Leaflet map · Copernicus DEM · OpenStreetMap"), ("Backend", "Python · FastAPI · pandas · SQLite"),
         ("Frontend", "React + TypeScript web app"), ("Hardware", "One on-prem laptop or server · IoT fuel sensors")]
for i, (cat, tech) in enumerate(stack):
    shp(s3, MSO_SHAPE.ROUNDED_RECTANGLE, 0.4 + i * 2.1, 1.58, 2.02, 1.0, fill=TINT, radius=0.12, margin=0.08,
        paras=[[(cat, {"bold": True, "color": NAVY, "size": 13})], *tech.split("\n")], size=12, color=INK)

label(s3, 0.4, 2.76, 12.5, "Methodology and process for implementation (Flow Charts/Images/ working prototype)")
flow = [
    (0.4, 2.8, "10 data streams", [[("Real: ", {"bold": True, "color": GREEN}), ("weather since 1995, 16-day forecast, terrain (DEM), highways (OSM)", {})],
                                    [("Simulated: ", {"bold": True, "color": AMBER}), ("CICG-like stock & issues, IoT fuel sensors, BRO roads, IAF airlift", {})]]),
    (3.55, 1.85, "Adapters", ["1 config file per stream", "Unknown codes refused", "Health checks"]),
    (5.75, 1.5, "One database", ["SQLite on the unit's own machine"]),
    (7.6, 3.05, "Engines", [[("Forecast", {"bold": True}), (" (LightGBM) → ", {}), ("Point of No Return", {"bold": True}),
                             (" (31 real winters) → ", {}), ("Mixed-fleet plan", {"bold": True}), (" (CP-SAT, 7 days)", {})],
                            "+ drone advisor · what-if · data alerts"]),
    (11.0, 1.95, "Commander", ["9 screens", "Approve / Modify / Reject", "Audit log"]),
]
for x, w, head, body in flow:
    paras = [{"runs": [(head, {"bold": True, "color": NAVY, "size": 13.5})], "after": 3}]
    paras += [{"runs": [(b, {})] if isinstance(b, str) else b, "after": 3} for b in body]
    shp(s3, MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.12, w, 1.95, fill=WHITE, line=NAVY, line_w=1.25, radius=0.08,
        margin=0.1, paras=paras, size=12, color=INK)
for x in (3.215, 5.415, 7.265, 10.665):
    shp(s3, MSO_SHAPE.RIGHT_ARROW, x, 3.9, 0.32, 0.4, fill=GREY)

pic(s3, "04-mixed-fleet-plan.png", 0.4, 5.3, 8.7, crop=(98, 181, 1352, 390))
tb(s3, 9.3, 5.3, 3.65, 1.45, [
    {"runs": [("Working prototype", {"bold": True, "color": NAVY, "size": 14})], "after": 4},
    {"runs": [("Every plan line explains itself: route, terrain clearance, sorties and cost vs helicopter, "
               "then Approve / Modify / Reject.", {})], "after": 4},
    [("Costs are PLACEHOLDER.", {"italic": True, "color": GREY, "size": 11})]], size=12)

# ----- Slide 4: FEASIBILITY AND VIABILITY -----
drop(get(s4, "TextBox 8"))
label(s4, 0.4, 1.22, 4.6, "Analysis of the feasibility of the idea")
feas = [("Built and running: ", "9 screens, one-click start, guided demo"),
        ("Works offline ", "on one laptop; no data leaves it"),
        ("Free, open-source stack: ", "no licence or cloud cost"),
        ("Plugs into existing systems: ", "reads CICG, BRO and IAF exports without changing them"),
        ("Stress-tested: ", "troops ×1.5 → posts at risk 3 → 6; pass closes 10 days early → helicopter sorties 30 → 49")]
y = 1.66
for lead, rest in feas:
    icon(s4, 0.42, y + 0.03, 0.34, "✓", GREEN)
    tb(s4, 0.88, y, 4.1, 0.9, [[(lead, {"bold": True, "color": NAVY}), (rest, {})]], size=13.5)
    y += 0.97 if len(rest) < 60 else 1.15

tb(s4, 5.25, 1.22, 7.7, 0.36, [[("Limitation → Mitigation", {"bold": True, "color": NAVY, "size": 14}),
                                  ("    our honest gaps, and how we handle them", {"italic": True, "color": GREY})]], size=12)
rows = [
    ("Military data is classified, so our stock and use data are simulated",
     "Adapters read real CICG, BRO and IAF exports through one config file each; simulated data is labelled on every screen"),
    ("Forecast accuracy is measured only on simulated data, built from the same weather inputs",
     "The model is used only where it beats a 28-day average on held-out days; real accuracy is measured once CICG data is connected"),
    ("Forward-track closure dates use PLACEHOLDER scaling of the pass rule",
     "Pass rules are fitted on real closure dates; the unit's own track records replace the placeholders"),
    ("Only 4–5 real closures per pass (error ±13 / ±18 days)",
     "Plan at 80% confidence, not the middle date; add BRO's records as they come"),
    ("The demo scenario is designed, and all costs are PLACEHOLDER",
     "The same engine was re-run with changed inputs and responded correctly; real costs go into one parameter file"),
    ("Large sectors may hit the solver's 20-second limit; roles are demo-only",
     "The screen then says \"best plan found within time limit\"; Army SSO/PKI supplies roles in deployment"),
]
gf = s4.shapes.add_table(len(rows) + 1, 2, Inches(5.25), Inches(1.64), Inches(7.7), Inches(4.9))
table = gf.table
table.columns[0].width, table.columns[1].width = Inches(3.35), Inches(4.35)
head = ("Limitation (potential challenges and risks)", "Mitigation (strategies for overcoming these challenges)")
for r_i in range(len(rows) + 1):
    table.rows[r_i].height = Inches(0.5 if r_i == 0 else 0.73)
    for c_i in range(2):
        cell = table.cell(r_i, c_i)
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r_i == 0:
            cell.fill.fore_color.rgb = NAVY
            cell.text_frame.text = ""
            fill_tf(cell.text_frame, [[(head[c_i], {"bold": True})]], size=12, color=WHITE)
        else:
            cell.fill.fore_color.rgb = TINT if r_i % 2 else WHITE
            cell.text_frame.text = ""
            text = rows[r_i - 1][c_i]
            fill_tf(cell.text_frame, [[(text, {"bold": c_i == 0, "color": RED if c_i == 0 else INK})]], size=11.5)
# the empty first paragraph left by text="" is reused by fill_tf, so no blank lines remain

# ----- Slide 5: IMPACT AND BENEFITS -----
drop(get(s5, "TextBox 8"))
label(s5, 0.4, 1.22, 4.6, "Potential impact on the target audience")
tb(s5, 0.4, 1.56, 4.6, 0.4, [[("For ", {}), ("commanders, logistics staff and post quartermasters", {"bold": True})]], size=13)
cards = [("29 Oct", NAVY, "Alpha's last safe road day, at 80% confidence, known before the pass closes", "simulated · PLACEHOLDER track rules"),
         ("30 → 0", MAGENTA, "Foxtrot's helicopter sorties this week, with 20 drones at the advised base", "simulated · PLACEHOLDER costs")]
for i, (big, col, text, tg) in enumerate(cards):
    y = 2.2 + i * 2.27
    shp(s5, MSO_SHAPE.ROUNDED_RECTANGLE, 0.4, y, 4.6, 2.12, fill=TINT, radius=0.08)
    tb(s5, 0.6, y + 0.12, 4.2, 0.8, [[(big, {"bold": True, "color": col})]], size=40)
    tb(s5, 0.6, y + 0.95, 4.2, 0.7, [text], size=13.5)
    tag(s5, 0.62, y + 1.68, 3.3 if "track" in tg else 2.75, tg)

label(s5, 5.35, 1.22, 7.6, "Benefits of the solution (social, economic, environmental, etc.)")
tiles = [("⏰", BLUE, "Decide before passes close", "Last safe road day for every post and supply"),
         ("⛑", RED, "Less risk to troops", "Fewer risky helicopter sorties and porter / mule trips"),
         ("✈", MAGENTA, "Ready for 2,715 drones", "Advises where to base them, by altitude band"),
         ("✓", GREEN, "Full accountability", "Every approve, modify or reject is logged"),
         ("⇄", TEAL, "Plugs into existing systems", "Reads CICG and other exports; replaces nothing"),
         ("⛰", NAVY, "Scales to other sectors", "Any high-altitude sector: new data, same engine")]
for i, (g, col, head_t, desc) in enumerate(tiles):
    x, y = 5.35 + (i % 2) * 3.9, 1.64 + (i // 2) * 1.62
    shp(s5, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.7, 1.46, fill=WHITE, line=EDGE, line_w=1, radius=0.1)
    icon(s5, x + 0.18, y + 0.2, 0.52, g, col)
    tb(s5, x + 0.85, y + 0.14, 2.75, 1.2, [{"runs": [(head_t, {"bold": True, "color": NAVY, "size": 14})], "after": 3}, desc], size=12.5)

# ----- Slide 6: RESEARCH AND REFERENCES -----
drop(get(s6, "TextBox 8"))
label(s6, 0.4, 1.22, 12.5, "Details / Links of the reference and research work")
L = {
    "idrw": "https://idrw.org/indian-army-floats-rfi-for-2715-logistic-drones-across-three-altitude-bands-for-last-mile-border-resupply/",
    "outlook": "https://www.outlookindia.com/national/defence-ministry-seeks-2715-logistic-drones-to-strengthen-indian-army-supply-chain-on-china-pakistan-borders-2",
    "tribune22": "https://www.tribuneindia.com/news/nation/specialised-logistics-drone-to-deliver-supplies-in-himalayas-442161",
    "primer": "https://www.globalsecurity.org/wmd/library/news/india/2025/india-250810-india-pib03.htm",
    "week": "https://www.theweek.in/theweek/current/2020/08/13/winter-is-coming.html",
    "sps": "https://www.spslandforces.com/experts-speak/?id=1339&h=Logistics-Drones-and-Animal-Transport",
    "qforce": "https://newsonair.gov.in/indian-army-launches-q-force-application-for-logistics-and-sustainment-management/",
    "gk": "https://www.greaterkashmir.com/kashmir/srinagar-leh-highway-zojila-pass-reopens-in-record-time/",
    "tribune73": "https://www.tribuneindia.com/news/j-k/srinagar-leh-road-reopens-after-73-days-378993",
    "print": "https://theprint.in/india/strategic-srinagar-leh-national-highway-reopened-for-traffic-in-record-time/1448166/",
    "air_zoji": "https://www.newsonair.gov.in/zojila-pass-nh1-srinagar-leh-highway-remains-open-during-february-for-the-first-time-ever",
    "air_ml25": "https://www.newsonair.gov.in/leh-manali-national-highway-closed-for-all-vehicles-due-to-winter-conditions/",
    "air_ml24": "https://www.newsonair.gov.in/bro-opened-leh-manali-national-highway-for-movement-of-army-vehicles",
    "hroutes": "https://www.himalayan-routes.com/manali-leh-highway-status-2023/",
    "om": "https://open-meteo.com/en/docs/historical-weather-api",
    "dem": "https://open-meteo.com/en/docs/elevation-api",
    "osm": "https://www.openstreetmap.org/copyright",
    "sabal": "https://thedefensepost.com/2024/12/05/indian-army-sabal-drones",
}


def ref(title, links):
    runs = []
    for k, (txt, key) in enumerate(links):
        if k:
            runs.append(("  ·  ", {"color": GREY}))
        runs.append((txt, {"color": BLUE, "underline": True, "link": L[key], "size": 11}))
    return [{"runs": [(title, {"bold": True})], "after": 1}, {"runs": runs, "after": 7}]


def group(head_t):
    return {"runs": [(head_t, {"bold": True, "color": NAVY, "size": 14})], "after": 5}


col1 = [group("Problem and policy")]
col1 += ref("MoD RFI: 2,715 logistic drones, 3 altitude bands (Aug 2026)", [("idrw.org", "idrw"), ("Outlook", "outlook")])
col1 += ref("Army RFP: 363 drones, 15 kg at high altitude, 10 km range", [("The Tribune, 17 Oct 2022", "tribune22")])
col1 += ref("Joint Primer for Integrated Logistics, HQ IDS (Aug 2025)", [("PIB via GlobalSecurity", "primer")])
col1 += ref("Winter stocking: ~2 lakh t for ~6 months; cost quotes", [("The Week, 13 Aug 2020", "week")])
col1 += ref("Mule units retained to ~2032", [("SP's Land Forces, Nov 2025", "sps")])
col1 += ref("Q-FORCE logistics app (May 2026)", [("AIR News, 22 May 2026", "qforce")])
def note(text):
    """PowerPoint's PDF export drops the link in a box's last reference unless a visible line follows it."""
    return [[(text, {"italic": True, "color": GREY, "size": 10.5})]]


tb(s6, 0.4, 1.62, 4.35, 5.2, col1 + note("Full list with confidence levels: docs/FACTS.md"), size=12)

col2 = [group("Real closure dates")]
col2 += ref("Zoji La closures and reopenings, 2020–25",
            [("Greater Kashmir", "gk"), ("The Tribune", "tribune73"), ("ThePrint", "print"), ("AIR News", "air_zoji")])
col2 += ref("Manali–Leh closures and openings, 2021–26",
            [("AIR News 2025", "air_ml25"), ("AIR News 2024", "air_ml24"), ("himalayan-routes", "hroutes")])
col2 += [group("Open data used (real, cached offline)")]
col2 += ref("Weather since 1995 and 16-day forecast", [("Open-Meteo ERA5 (CC BY 4.0)", "om")])
col2 += ref("Terrain heights and national highways", [("Copernicus DEM GLO-90", "dem"), ("OpenStreetMap (ODbL)", "osm")])
col2 += ref("Logistics drone specs: Sabal 20", [("The Defense Post, Dec 2024", "sabal")])
tb(s6, 4.95, 1.62, 4.35, 5.2, col2 + note("All data is free and public; no account or key needed"), size=12)

shp(s6, MSO_SHAPE.ROUNDED_RECTANGLE, 9.55, 1.62, 3.4, 5.05, fill=TINT, radius=0.06)
tb(s6, 9.75, 1.75, 3.0, 0.4, [[("What we verified", {"bold": True, "color": NAVY})]], size=15)
checks = [("Real data: ", "31 winters of weather, terrain, highways and closure dates"),
          ("Prior work: ", "3 public prototypes for this problem studied; none plans drones or gives a closure probability"),
          ("Self-check: ", "outputs change correctly when inputs change. Military data simulated; costs PLACEHOLDER")]
y = 2.3
for lead, rest in checks:
    icon(s6, 9.75, y + 0.04, 0.3, "✓", GREEN)
    tb(s6, 10.15, y, 2.7, 1.3, [[(lead, {"bold": True, "color": NAVY}), (rest, {})]], size=12.5)
    y += 1.42

prs.save(str(OUT))
print("saved", OUT)
